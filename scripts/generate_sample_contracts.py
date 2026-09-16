"""Generate sample contracts for trying the review by hand.

The fixtures under tests/fixtures/ exist to be asserted against and are all plain
text. These are for driving the product: PDF and DOCX as a user would actually
upload them, a contract whose dangerous term lives only inside a table, one with
no clause numbering at all, and a scan with no text layer that must be refused
rather than reviewed.

Generated rather than committed as binaries so the source wording stays readable
and reviewable, in the same spirit as the PDF fixture built inside the test suite.

    uv run python -m scripts.generate_sample_contracts
"""

import io
import pathlib

OUT = pathlib.Path("samples")

# --------------------------------------------------------------------------
# 01 - Employment, heavily one-sided. The full HOT path, in PDF.
# --------------------------------------------------------------------------
EMPLOYMENT_LOADED = """EMPLOYMENT AGREEMENT

This Employment Agreement is made at Bengaluru on 12 August 2025 between
Vantage Systems India Private Limited, a company incorporated under the
Companies Act, 2013, having its registered office at Whitefield, Bengaluru
(hereinafter "the Company"), and Ms. K. Raghavan, residing at HSR Layout,
Bengaluru (hereinafter "the Employee").

1. APPOINTMENT AND TERM
The Company appoints the Employee as Product Manager with effect from
1 September 2025. This Agreement shall remain in force until terminated in
accordance with the provisions set out below.

2. REMUNERATION
The Employee shall be paid a gross salary of Rs. 24,00,000 per annum, payable
monthly in arrears on or before the seventh day of the succeeding month,
subject to deduction of tax at source and other statutory deductions.

3. PROBATION
The Employee shall be on probation for a period of six months from the date of
joining. The Company may extend the probation period at its sole discretion.

4. MINIMUM SERVICE PERIOD
The Employee shall serve the Company for a minimum service period of three
years from the date of joining. In the event the Employee resigns before the
completion of the said period, the Employee shall refund the training cost
incurred by the Company amounting to Rs. 3,00,000 together with all recruitment
expenses incurred in connection with the Employee's appointment.

5. TERMINATION
The Company may terminate this Agreement at any time without assigning any
reason by giving thirty days notice in writing. The Employee may resign only
after serving a notice period of ninety days, failing which the Employee shall
be liable to pay salary in lieu of the shortfall.

6. CONFIDENTIALITY
The Employee shall not disclose any Confidential Information of the Company to
any third party. The obligations of confidentiality under this clause shall
survive the termination of this Agreement in perpetuity and shall remain
binding upon the Employee for all time.

7. INTELLECTUAL PROPERTY
The Employee hereby assigns to the Company all intellectual property,
inventions, designs and works created by the Employee during the term of
employment, whether or not such creation is related to the business of the
Company and whether or not created during working hours or using the resources
of the Company.

8. NON-COMPETITION
The Employee shall not, for a period of twenty four months after the
termination of employment with the Company, directly or indirectly engage in,
be employed by, or render services to any competing business anywhere in India.

9. NON-SOLICITATION
The Employee shall not solicit any employee or customer of the Company for a
period of eighteen months following cessation of employment.

10. DISPUTE RESOLUTION
Any dispute arising out of or in connection with this Agreement shall be
referred to arbitration before a sole arbitrator appointed by the Company,
whose decision shall be final and binding. The seat of arbitration shall be
Bengaluru.

11. GOVERNING LAW
This Agreement shall be governed by the laws of India and the courts at
Bengaluru shall have exclusive jurisdiction.
"""

# --------------------------------------------------------------------------
# 04 - Employment, genuinely fair. The false-positive check.
# --------------------------------------------------------------------------
EMPLOYMENT_FAIR = """EMPLOYMENT AGREEMENT

This Employment Agreement is made at Pune on 3 June 2025 between Meridian
Analytics Private Limited, having its registered office at Baner, Pune
(hereinafter "the Company"), and Mr. S. Kulkarni, residing at Aundh, Pune
(hereinafter "the Employee").

1. APPOINTMENT AND TERM
The Company appoints the Employee as Data Engineer with effect from 16 June
2025. This Agreement shall continue until brought to an end in accordance with
clause 6 below.

2. REMUNERATION
The Employee shall be paid a gross salary of Rs. 18,00,000 per annum, payable
monthly in arrears on or before the last working day of each month, subject to
statutory deductions.

3. PROBATION
The Employee shall be on probation for three months from the date of joining,
on completion of which the Employee shall be confirmed in writing.

4. LEAVE
The Employee shall be entitled to twenty days of earned leave and twelve days
of casual and sick leave in each calendar year, in addition to public holidays
declared by the Company.

5. CONFIDENTIALITY
The Employee shall keep confidential the trade secrets and customer information
of the Company. This obligation shall continue for two years after the Employee
leaves the Company, and shall not apply to information that is or becomes
public through no fault of the Employee, or that the Employee is required by
law to disclose.

6. TERMINATION
Either party may bring this Agreement to an end by giving the other sixty days
written notice, or salary in lieu of that notice.

7. INTELLECTUAL PROPERTY
Work product created by the Employee in the course of the Employee's duties and
using the resources of the Company shall belong to the Company. Anything the
Employee creates on the Employee's own time, without the resources of the
Company and unrelated to the business of the Company, remains the property of
the Employee.

8. DISPUTE RESOLUTION
Any dispute shall be referred to a sole arbitrator appointed by mutual written
agreement of the parties, failing which by the Mumbai Centre for International
Arbitration. The seat of arbitration shall be Pune.

9. GOVERNING LAW
This Agreement shall be governed by the laws of India and the courts at Pune
shall have jurisdiction.
"""

# --------------------------------------------------------------------------
# 03 - Freelance, plain text.
# --------------------------------------------------------------------------
FREELANCE = """INDEPENDENT CONTRACTOR AGREEMENT

This Agreement is made on 14 March 2025 between Zephyr Media Private Limited,
Mumbai (hereinafter "the Client"), and Ms. A. Bose, a freelance designer
residing at Salt Lake, Kolkata (hereinafter "the Contractor").

1. SERVICES
The Contractor shall provide graphic design services as described in each
statement of work agreed between the parties from time to time.

2. FEES
The Client shall pay the Contractor Rs. 80,000 per month within forty five days
of receipt of a valid invoice for the preceding month.

3. INTELLECTUAL PROPERTY
The Contractor hereby assigns to the Client all intellectual property, designs,
artwork and materials created by the Contractor during the term of this
Agreement, whether or not created for the Client and whether or not created
using the resources of the Client.

4. TERMINATION
The Client may terminate this Agreement at any time without assigning any
reason by written notice with immediate effect. The Contractor may terminate
only by giving ninety days written notice.

5. DELAY
In the event of delay in delivery of any deliverable, the Contractor shall pay
a penalty of Rs. 2,000 per day of delay until the deliverable is accepted.

6. INDEMNITY
The Contractor shall indemnify, defend and hold harmless the Client against any
and all claims, losses, damages and expenses of whatever nature arising out of
or in connection with the deliverables or this Agreement.

7. NON-COMPETITION
The Contractor shall not, for a period of twelve months after the termination
of this Agreement, provide design services to any competing business of the
Client.

8. GOVERNING LAW
This Agreement shall be governed by the laws of India and the courts at Mumbai
shall have exclusive jurisdiction.
"""

# --------------------------------------------------------------------------
# 05 - An offer letter with no clause numbering at all.
# --------------------------------------------------------------------------
OFFER_LETTER = """Harbour Labs Private Limited
4th Floor, Prestige Tower, Koramangala, Bengaluru 560034

7 October 2025

Dear Mr. Menon,

We are delighted to offer you the position of Senior Backend Engineer at
Harbour Labs Private Limited. This letter sets out the terms on which we make
that offer, and we would be grateful if you would sign and return a copy to
confirm your acceptance.

Your annual cost to company will be Rs. 32,00,000, made up of a fixed component
of Rs. 27,00,000 and a performance bonus of up to Rs. 5,00,000 payable at the
sole discretion of the management after the close of each financial year. Your
salary will be credited on the last working day of each calendar month after
deduction of tax at source and any other statutory deductions.

Your employment will commence on 3 November 2025 and will be subject to a
probationary period of six months, during which either party may end the
engagement by giving seven days notice. On confirmation, the notice period
required from you will be ninety days, while the Company may end your
employment by giving you thirty days notice or salary in lieu thereof.

During your employment and for a period of eighteen months after you leave the
Company for any reason, you shall not accept employment with, or render
services to, any organisation that competes with the business of the Company
anywhere in India, nor shall you solicit any employee or client of the Company.

All inventions, designs, source code and other works that you create during
your employment shall vest absolutely in the Company, whether or not they were
created in the course of your duties and whether or not the resources of the
Company were used in creating them.

You shall hold in strict confidence all information belonging to the Company
that comes to your knowledge, and this obligation shall continue in perpetuity
after the termination of your employment.

Any dispute arising out of your employment shall be referred to arbitration
before a sole arbitrator nominated by the Company, and the seat of arbitration
shall be Bengaluru.

We look forward to welcoming you to the team.

Yours sincerely,

R. Deshpande
Head of People Operations
Harbour Labs Private Limited
"""

# --------------------------------------------------------------------------
# 02 - Rental. The deposit terms exist ONLY inside a table.
# --------------------------------------------------------------------------
RENTAL_PARAGRAPHS = [
    "LEAVE AND LICENCE AGREEMENT",
    "This Agreement is executed at Pune on 5 March 2025 between Mrs. R. Deshpande, "
    "residing at Baner, Pune (hereinafter \"the Licensor\"), and Mr. K. Iyer, "
    "residing at Kothrud, Pune (hereinafter \"the Licensee\").",
    "1. PREMISES AND TERM",
    "The Licensor grants the Licensee a licence to occupy Flat No. 402, Sunrise "
    "Residency, Baner, Pune for a period of eleven months with effect from "
    "15 March 2025.",
    "2. LICENCE FEE",
    "The Licensee shall pay the monthly licence fee stated in the Schedule below "
    "on or before the fifth day of each calendar month, without any demand being "
    "made by the Licensor.",
    "3. DEPOSIT",
    "The Licensee shall pay the security deposit stated in the Schedule below on "
    "or before the commencement of the licence period. The terms governing the "
    "deposit are those set out in that Schedule.",
    "4. LOCK-IN PERIOD",
    "The Licensee shall not vacate the premises before the expiry of nine months "
    "from the commencement date. This lock-in period is binding on the Licensee "
    "notwithstanding any circumstances whatsoever.",
    "5. DELAY IN PAYMENT",
    "In the event of delay in payment of the licence fee, the Licensee shall pay "
    "a penalty of Rs. 1,000 per day of delay until the date of actual payment.",
    "6. TERMINATION",
    "The Licensor may terminate this Agreement at any time without assigning any "
    "reason by giving thirty days notice. The Licensee shall vacate the premises "
    "upon receipt of such notice.",
    "7. GOVERNING LAW",
    "This Agreement shall be governed by the laws of India and the courts at Pune "
    "shall have exclusive jurisdiction over any dispute.",
    "SCHEDULE",
]

# The dangerous terms live here and nowhere else. document.paragraphs skips
# tables entirely, so this is what a table-blind parser silently loses.
RENTAL_TABLE = [
    ["Particulars", "Terms"],
    ["Monthly licence fee", "Rs. 38,000 per month, escalating by ten per cent on each renewal"],
    [
        "Security deposit",
        "An interest-free security deposit equivalent to ten months licence fee, "
        "that is Rs. 3,80,000, which shall stand forfeited in favour of the "
        "Licensor if the Licensee vacates before the expiry of the lock-in period",
    ],
    ["Maintenance", "Payable by the Licensee in addition to the licence fee"],
]

# --------------------------------------------------------------------------
# 08 - Loan. Unilateral interest variation and an ouster of remedy.
# --------------------------------------------------------------------------
LOAN_PARAGRAPHS = [
    "LOAN AGREEMENT",
    "This Loan Agreement is made on 2 May 2025 between Suvarna Credit Private "
    "Limited, having its registered office at Nariman Point, Mumbai (hereinafter "
    "\"the Lender\"), and Mr. D. Pillai, residing at Thane (hereinafter \"the "
    "Borrower\").",
    "1. LOAN",
    "The Lender agrees to lend to the Borrower the principal amount of "
    "Rs. 5,00,000 on the terms set out in this Agreement.",
    "2. INTEREST",
    "The Loan shall carry interest at the rate of eighteen per cent per annum. "
    "The Lender reserves the right to amend the rate of interest at any time at "
    "its sole discretion by posting the revised rate on its website.",
    "3. REPAYMENT",
    "The Borrower shall repay the Loan together with interest in thirty six "
    "equated monthly instalments of Rs. 18,076 each, payable on the fifth day of "
    "every month.",
    "4. DEFAULT",
    "In the event of delay in payment of any instalment, the Borrower shall pay a "
    "penalty of Rs. 1,500 per day of delay in addition to interest on the overdue "
    "amount.",
    "5. DISPUTE RESOLUTION",
    "Any dispute arising out of this Agreement shall be referred to arbitration "
    "before a sole arbitrator appointed by the Lender, whose decision shall be "
    "final and binding on the Borrower.",
    "6. WAIVER",
    "The Borrower shall not be entitled to institute any suit or proceeding "
    "against the Lender in any court in respect of this Agreement, and hereby "
    "waives all statutory rights and remedies available to a borrower under "
    "applicable law.",
    "7. GOVERNING LAW",
    "This Agreement shall be governed by the laws of India and the courts at "
    "Mumbai shall have exclusive jurisdiction.",
]


def _msa() -> str:
    """A long master services agreement, to exercise the triage caps.

    Also carries a reference to a Schedule that is not attached and a
    cross-reference to a clause that does not exist.
    """
    head = """MASTER SERVICES AGREEMENT

This Master Services Agreement is entered into on 12 February 2025 between
Helioscope Cloud Services Private Limited, Hyderabad (hereinafter "the
Provider"), and Bluefin Retail Private Limited, Chennai (hereinafter "the
Client").

1. SERVICES
The Provider shall make available to the Client the software-as-a-service
platform described in Schedule A, together with such support services as are
specified in that Schedule.

2. FEES AND PAYMENT
The Client shall pay the subscription fees set out in the Order Form within
thirty days of the date of each invoice. Fees are exclusive of applicable taxes.

3. TERM AND RENEWAL
The initial subscription term is twelve months. This Agreement shall be
automatically renewed for successive periods of twelve months unless either
party gives written notice of non-renewal not less than ninety days prior to
the end of the then current term.

4. MODIFICATION OF TERMS
The Provider reserves the right to amend these terms at any time at its sole
discretion by posting the revised terms to its website. Continued use of the
platform after such posting shall constitute acceptance of the revised terms.

5. INDEMNITY
The Client shall indemnify, defend and hold harmless the Provider against any
and all claims, losses, damages, liabilities, costs and expenses of whatever
nature arising out of or in connection with the use of the platform by the
Client.

6. PERSONAL DATA
The Client consents to the processing of personal data uploaded to the
platform. The Provider may share such personal data with its affiliates,
sub-processors and third party analytics partners for the purpose of improving
the platform.

7. ASSIGNMENT
The Client shall not assign or transfer this Agreement without the prior
written consent of the Provider. The Provider may assign this Agreement to any
affiliate or successor entity without restriction and without notice to the
Client.

8. LIMITATION OF REMEDIES
The Client hereby waives all statutory rights and remedies available to it in
respect of any defect in the platform, to the fullest extent permitted.

9. GOVERNING LAW AND JURISDICTION
This Agreement shall be governed by the laws of Singapore and the courts at
Singapore shall have exclusive jurisdiction over any dispute arising hereunder.

10. TERMINATION
The Provider may terminate this Agreement at any time without assigning any
reason by giving thirty days written notice to the Client, and the consequences
set out in clause 47 shall then apply to any fees already paid.
"""
    # Ordinary operational clauses, to push the clause count past the HOT cap.
    filler = [
        ("SERVICE LEVELS", "The Provider shall use reasonable endeavours to make the platform available "
                           "for 99.5 per cent of each calendar month, excluding scheduled maintenance "
                           "notified to the Client in advance."),
        ("SUPPORT", "The Provider shall provide support by electronic mail between 0930 and 1830 "
                    "hours on business days, and shall acknowledge each request within one "
                    "business day of its receipt."),
        ("SCHEDULED MAINTENANCE", "The Provider may carry out scheduled maintenance outside business hours "
                                  "on giving the Client not less than five business days notice by "
                                  "electronic mail to the address notified for that purpose."),
        ("ACCOUNT ADMINISTRATION", "The Client shall nominate an administrator who shall be responsible for "
                                   "managing user accounts, and shall notify the Provider in writing of any "
                                   "change in the identity of that administrator."),
        ("ACCEPTABLE USE", "The Client shall not use the platform to store or transmit any material "
                           "that is unlawful, and shall ensure that its users comply with the "
                           "acceptable use policy notified to the Client from time to time."),
        ("USER CREDENTIALS", "The Client shall keep all user credentials confidential and shall notify "
                             "the Provider without delay on becoming aware of any unauthorised use of "
                             "any credential issued under this Agreement."),
        ("DATA BACKUP", "The Provider shall take backups of Client data at intervals of not more "
                        "than twenty four hours and shall retain each backup for a period of "
                        "thirty days from the date on which it was taken."),
        ("DATA EXPORT", "On written request made during the term, the Provider shall make available "
                        "to the Client an export of its data in a commonly used machine readable "
                        "format within fifteen business days of the request."),
        ("SECURITY", "The Provider shall maintain administrative, physical and technical "
                     "safeguards designed to protect the security and integrity of Client data "
                     "stored on the platform."),
        ("AUDIT", "The Provider shall on written request furnish to the Client a copy of its "
                  "most recent third party security assessment report, subject to the Client "
                  "first executing a confidentiality undertaking."),
        ("SUBCONTRACTING", "The Provider may engage subcontractors in the performance of the services "
                           "and shall remain responsible for the acts and omissions of any "
                           "subcontractor so engaged."),
        ("TRAINING", "The Provider shall provide up to eight hours of remote onboarding training "
                     "for the users of the Client within sixty days of the commencement date of "
                     "this Agreement."),
        ("PUBLICITY", "Neither party shall issue any press release referring to the other party "
                      "without first obtaining the written approval of that party to the form and "
                      "content of the release."),
        ("NOTICES", "Any notice under this Agreement shall be given in writing and shall be sent "
                    "to the address of the recipient stated above, or to such other address as "
                    "that party may notify in writing."),
        ("FORCE MAJEURE", "Neither party shall be liable for any failure to perform its obligations "
                          "to the extent that the failure is caused by an event beyond its "
                          "reasonable control, provided it notifies the other party promptly."),
        ("SEVERABILITY", "If any provision of this Agreement is held to be invalid or unenforceable, "
                         "the remaining provisions shall continue in full force and effect as "
                         "though the invalid provision had not been included."),
        ("WAIVER", "No failure or delay by either party in exercising any right under this "
                   "Agreement shall operate as a waiver of that right, nor shall any single "
                   "exercise of it preclude any further exercise."),
        ("ENTIRE AGREEMENT", "This Agreement, together with the Order Form, constitutes the entire "
                             "agreement between the parties and supersedes all prior discussions and "
                             "understandings relating to its subject matter."),
        ("SURVIVAL", "The provisions of this Agreement which by their nature are intended to "
                     "survive its termination shall continue in effect after termination for so "
                     "long as is necessary to give them effect."),
        ("COUNTERPARTS", "This Agreement may be executed in counterparts, each of which shall be "
                         "deemed an original, and all of which together shall constitute one and "
                         "the same instrument."),
    ]
    parts = [head]
    for index, (heading, body) in enumerate(filler, start=11):
        parts.append(f"{index}. {heading}\n{body}\n")
    return "\n".join(parts)


# --------------------------------------------------------------------------
# Writers
# --------------------------------------------------------------------------

def write_txt(name: str, text: str) -> pathlib.Path:
    path = OUT / name
    path.write_text(text, encoding="utf-8")
    return path


def write_pdf(name: str, text: str, title: str) -> pathlib.Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    from src.core.report_export import BODY_FONT

    path = OUT / name
    document = SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=22 * mm, rightMargin=22 * mm,
        topMargin=20 * mm, bottomMargin=20 * mm, title=title,
    )
    body = ParagraphStyle(
        "body", parent=getSampleStyleSheet()["Normal"],
        fontName=BODY_FONT, fontSize=10, leading=14,
    )
    story = []
    for block in text.split("\n\n"):
        if not block.strip():
            continue
        # <br/> keeps the line wrapping a real contract PDF has, which is what
        # exercises the whitespace-tolerant matching.
        story.append(Paragraph(block.strip().replace("\n", "<br/>"), body))
        story.append(Spacer(1, 7))
    document.build(story)
    return path


def write_docx(name: str, paragraphs, table_rows=None) -> pathlib.Path:
    import docx

    path = OUT / name
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        table.style = "Table Grid"
        for row_index, row in enumerate(table_rows):
            for cell_index, value in enumerate(row):
                table.cell(row_index, cell_index).text = value
    document.save(str(path))
    return path


def write_scan_of(name: str, source: pathlib.Path, dpi: int = 200,
                  rotate: float = 0.4, noise: int = 10) -> pathlib.Path:
    """Turn a text-layer PDF into a genuine image-only one.

    Rendered to images and written back with no text layer, then degraded the
    way a real scan is: a fraction of a degree of skew from a page fed slightly
    crooked, sensor noise, and a resolution below the 300 DPI OCR prefers. A
    pristine render would flatter the OCR path and prove nothing.
    """
    import numpy as np
    import pypdfium2 as pdfium
    from PIL import Image

    document = pdfium.PdfDocument(str(source))
    pages = []
    for index in range(len(document)):
        image = document[index].render(scale=dpi / 72).to_pil().convert("L")
        if rotate:
            image = image.rotate(rotate, resample=Image.BICUBIC, fillcolor=255, expand=False)
        if noise:
            array = np.asarray(image).astype(np.int16)
            rng = np.random.default_rng(index)
            array = np.clip(array + rng.integers(-noise, noise + 1, array.shape), 0, 255)
            image = Image.fromarray(array.astype(np.uint8))
        pages.append(image.convert("RGB"))

    path = OUT / name
    pages[0].save(str(path), format="PDF", save_all=True, append_images=pages[1:])
    return path


def write_scanned_pdf(name: str, pages: int = 3) -> pathlib.Path:
    """A PDF with no text layer, as a phone photo or a flatbed scan produces.

    Must be refused. Reviewing it would report no problems simply because
    nothing could be read, which reads exactly like a clean contract.
    """
    from pypdf import PdfWriter

    path = OUT / name
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=595, height=842)
    buffer = io.BytesIO()
    writer.write(buffer)
    path.write_bytes(buffer.getvalue())
    return path


def write_image_of(name: str, source: pathlib.Path, dpi: int = 170) -> pathlib.Path:
    """A photograph of the first page, as someone would send from a phone."""
    import numpy as np
    import pypdfium2 as pdfium
    from PIL import Image, ImageEnhance

    document = pdfium.PdfDocument(str(source))
    image = document[0].render(scale=dpi / 72).to_pil().convert("L")
    image = image.rotate(-0.8, resample=Image.BICUBIC, fillcolor=255, expand=False)
    # Uneven lighting, as a hand-held photo has.
    array = np.asarray(image).astype(np.float32)
    gradient = np.linspace(0.82, 1.06, array.shape[1])[None, :]
    array = np.clip(array * gradient, 0, 255)
    image = ImageEnhance.Contrast(Image.fromarray(array.astype(np.uint8))).enhance(0.92)

    path = OUT / name
    image.convert("RGB").save(str(path), format="PNG")
    return path


SAMPLES = [
    ("01_employment_loaded.pdf", "EMPLOYMENT", "EMPLOYEE",
     "Heavily one-sided employment contract. The full expensive path in PDF."),
    ("02_rental_table_terms.docx", "LEASE", "TENANT",
     "Deposit and forfeiture terms exist ONLY inside a table."),
    ("03_freelance_contract.txt", "FREELANCE", "SERVICE_PROVIDER",
     "Plain text baseline: uncapped indemnity, one-sided termination."),
    ("04_employment_fair.pdf", "EMPLOYMENT", "EMPLOYEE",
     "A genuinely fair contract. Should raise little or nothing."),
    ("05_offer_letter_unnumbered.txt", "EMPLOYMENT", "EMPLOYEE",
     "No clause numbering at all: forces the paragraph fallback."),
    ("06_scanned_no_text_layer.pdf", "EMPLOYMENT", "EMPLOYEE",
     "No text layer. Must be REFUSED, not reviewed."),
    ("07_master_services_long.pdf", "SAAS", "CLIENT",
     "30 clauses: exercises the HOT/WARM/COLD caps, a missing Schedule A and a broken cross-reference."),
    ("08_loan_agreement.docx", "LOAN", "BORROWER",
     "Unilateral interest variation, ouster of remedy, waiver of statutory rights."),
    ("09_scanned_employment.pdf", "EMPLOYMENT", "EMPLOYEE",
     "Contract 01 with no text layer: skewed, noisy, 200 DPI. Read by OCR."),
    ("10_photo_of_contract.png", "EMPLOYMENT", "EMPLOYEE",
     "A phone photograph of page 1: skewed, unevenly lit. Read by OCR."),
]


def main():
    OUT.mkdir(exist_ok=True)

    write_pdf("01_employment_loaded.pdf", EMPLOYMENT_LOADED, "Employment Agreement")
    write_docx("02_rental_table_terms.docx", RENTAL_PARAGRAPHS, RENTAL_TABLE)
    write_txt("03_freelance_contract.txt", FREELANCE)
    write_pdf("04_employment_fair.pdf", EMPLOYMENT_FAIR, "Employment Agreement")
    write_txt("05_offer_letter_unnumbered.txt", OFFER_LETTER)
    write_scanned_pdf("06_scanned_no_text_layer.pdf")
    write_pdf("07_master_services_long.pdf", _msa(), "Master Services Agreement")
    write_docx("08_loan_agreement.docx", LOAN_PARAGRAPHS)
    # A real scan: the same contract as 01, with no text layer at all.
    write_scan_of("09_scanned_employment.pdf", OUT / "01_employment_loaded.pdf")
    write_image_of("10_photo_of_contract.png", OUT / "01_employment_loaded.pdf")

    print(f"[+] Wrote {len(SAMPLES)} sample contracts to {OUT}/\n")
    for name, contract_type, position, note in SAMPLES:
        size = (OUT / name).stat().st_size
        print(f"  {name:34s} {contract_type:12s} as {position:16s} {size:7,d} bytes")
        print(f"  {'':34s} {note}")


if __name__ == "__main__":
    main()
