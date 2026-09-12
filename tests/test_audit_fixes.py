"""Regressions for the findings the phase 1-8 audit raised.

Each one was reproduced against the code before being fixed. What they have in
common is the failure mode: a clause, a term or a whole page went missing
silently, and an absent finding reads exactly like a clean contract.
"""

import asyncio
import io
from unittest.mock import patch

import docx
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from src.core.clause_segmenter import segment_clauses
from src.core.document_parser import DocumentParseError, normalize, parse_document
from src.core.red_flag_rules import evaluate_clauses, verify_quotes
from src.schemas.contract import PartyPosition

FILLER = (
    "\n\n5. NOTICES\nAny notice under this agreement shall be given in writing and "
    "sent to the address of the recipient stated above or as later notified."
    "\n\n6. SEVERABILITY\nIf any provision is held invalid the remaining provisions "
    "continue in full force as though it had never been included."
    "\n\n7. WAIVER\nNo delay by either party in exercising a right operates as a "
    "waiver of it, nor precludes any further exercise of that right."
)


def seg(text, name="c.txt"):
    document = parse_document(text.encode(), name)
    return document, segment_clauses(document)


def fired(clauses, position=PartyPosition.EMPLOYEE):
    return {f.rule_id for f in evaluate_clauses(clauses, position)}


# ---------------------------------------------------------------------------
# Heading detection
# ---------------------------------------------------------------------------

def test_an_abbreviation_does_not_end_the_heading():
    """"Rs." made the heading "The Lessee shall pay a penalty of Rs" and left
    "1,000 per day" as the body, so PENALTY_STIPULATION -- which needs both
    halves -- matched nothing at all."""
    text = (
        "3. The Lessee shall pay a penalty of Rs. 1,000 per day of delay in "
        "payment of rent until the date of actual payment of the amount due.\n\n"
        "4. The Lessee shall keep the premises in good and habitable condition "
        "throughout the currency of this agreement.\n\n"
        "5. This agreement is governed by the laws of India and the courts at "
        "Pune have exclusive jurisdiction over any dispute."
    )
    _, clauses = seg(text)
    assert clauses[0].heading is None
    assert "PENALTY_STIPULATION" in fired(clauses, PartyPosition.TENANT)


def test_a_short_opening_sentence_is_not_a_heading():
    text = (
        "3.1 The Company may terminate this Agreement forthwith. Such termination "
        "shall be without any liability to the Employee whatsoever.\n\n"
        "3.2 The Employee shall return all property of the Company upon such "
        "termination taking effect under this clause.\n\n"
        "3.3 This clause survives termination for a period of twelve months from "
        "the date on which it takes effect."
    )
    _, clauses = seg(text)
    assert clauses[0].heading is None


@pytest.mark.parametrize("caption", ["TERMINATION", "Termination Of Employment", "Notice Period"])
def test_a_real_caption_is_still_recognised(caption):
    text = (
        f"1. {caption}. The Company may terminate this Agreement at any time "
        "without assigning any reason by giving thirty days written notice.\n\n"
        "2. FEES\nThe Client shall pay each invoice within thirty days of its "
        "date without deduction or set-off of any kind.\n\n"
        "3. TERM\nThis agreement runs for twelve months from the commencement "
        "date stated above unless ended earlier."
    )
    _, clauses = seg(text)
    assert clauses[0].heading == caption


# ---------------------------------------------------------------------------
# Marker handling
# ---------------------------------------------------------------------------

def test_a_gap_in_the_numbering_does_not_lose_every_later_clause():
    """A clause deleted in negotiation leaves 1,2,3,7,8... The rejected marker
    never advanced the counter, so it poisoned all its successors and the whole
    contract fell back to paragraphs."""
    text = "".join(
        f"{n}. CLAUSE {n}\nThe parties agree that this clause {n} binds them and "
        "their successors in title without exception of any kind.\n\n"
        for n in [1, 2, 3, 7, 8, 9, 10]
    )
    _, clauses = seg(text)
    assert [c.number for c in clauses if c.number] == ["1", "2", "3", "7", "8", "9", "10"]


def test_a_spurious_number_inside_a_clause_is_still_rejected():
    """The gap tolerance must not readmit the noise it was guarding against."""
    text = (
        "1. PAYMENT\nThe fee is payable within thirty days of the invoice date by "
        "electronic transfer to the nominated account.\n"
        "2. NOTICE\nEither party may terminate on thirty days written notice, and "
        "9. of the Schedule then applies to amounts outstanding.\n"
        "3. GOVERNING LAW\nThis agreement is governed by the laws of India and the "
        "courts at Mumbai have jurisdiction."
    )
    _, clauses = seg(text)
    assert [c.number for c in clauses if c.number] == ["1", "2", "3"]


def test_an_article_style_marker_yields_its_heading():
    """"Article 1 APPOINTMENT" left "1 APPOINTMENT" in the body with no heading,
    because the marker was re-derived instead of remembered."""
    text = (
        "Article 1 APPOINTMENT\nThe Company appoints the Employee to the post "
        "named above with effect from the date in the schedule.\n\n"
        "Article 2 REMUNERATION\nThe Employee is paid the salary stated in the "
        "schedule, monthly in arrears after statutory deductions.\n\n"
        "Article 3 TERMINATION\nThe Company may terminate this Agreement at any "
        "time without assigning any reason on thirty days notice."
    )
    document, clauses = seg(text)
    assert [c.heading for c in clauses] == ["APPOINTMENT", "REMUNERATION", "TERMINATION"]
    for clause in clauses:
        body = document.text[clause.body_offset : clause.end_offset]
        assert not body.lstrip().startswith(clause.number)


def test_the_preamble_is_reviewed():
    """The title, recitals and definition of the parties were discarded, so a
    waiver buried in the recitals reached no rule at all."""
    text = (
        "EMPLOYMENT AGREEMENT\n\nThis Agreement is made between the Company and "
        "the Employee. The Employee hereby waives all statutory rights and "
        "remedies available under applicable law in respect of this engagement.\n\n"
        "1. SALARY\nThe Employee is paid the sum stated in the schedule monthly "
        "in arrears subject to deductions at source.\n\n"
        "2. TERM\nThis Agreement commences on the date above and continues until "
        "determined in accordance with its terms.\n\n"
        "3. NOTICE\nEither party may terminate on sixty days written notice to "
        "the other at the address stated above."
    )
    document, clauses = seg(text)
    assert clauses[0].start_offset == 0
    assert "UNFAIR_TERM_WAIVER" in fired(clauses)
    assert verify_quotes(evaluate_clauses(clauses, PartyPosition.EMPLOYEE), document.text) == []


def test_body_offset_is_right_when_the_body_repeats_the_heading():
    text = (
        "1. The Company may terminate this agreement\nThe Company may terminate "
        "this agreement at any time without assigning any reason on notice.\n\n"
        "2. TERM\nThis Agreement commences on the date above and continues until "
        "determined in accordance with its terms.\n\n"
        "3. NOTICE\nEither party may give notice to the other at the address "
        "stated above in accordance with this clause."
    )
    document, clauses = seg(text)
    body = document.text[clauses[0].body_offset : clauses[0].end_offset]
    assert body.startswith("The Company may terminate this agreement at any time")


# ---------------------------------------------------------------------------
# DOCX extraction
# ---------------------------------------------------------------------------

def _docx(build):
    document = docx.Document()
    build(document)
    for heading, body in [
        ("5. NOTICES", "Any notice shall be in writing and sent to the address stated above or as later notified in writing."),
        ("6. SEVERABILITY", "If any provision is held invalid the remainder continues in force as though it had never been included."),
        ("7. WAIVER", "No delay in exercising a right operates as a waiver of it nor precludes any further exercise of that right."),
    ]:
        document.add_paragraph(heading)
        document.add_paragraph(body)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _base(document):
    document.add_paragraph("1. FEES")
    document.add_paragraph("The Client shall pay each invoice within thirty days of its date without deduction.")
    document.add_paragraph("2. TERM")
    document.add_paragraph("This agreement runs for twelve months from the commencement date stated above.")


def test_a_clause_inside_a_content_control_is_read():
    """Everything inside a Word content control vanished, a whole clause with
    it, and nothing was said about the loss."""
    def build(document):
        _base(document)
        document.element.body.append(parse_xml(
            f'<w:sdt {nsdecls("w")}><w:sdtContent>'
            '<w:p><w:r><w:t>3. TERMINATION</w:t></w:r></w:p>'
            '<w:p><w:r><w:t>The Company may terminate this Agreement at any time '
            'without assigning any reason by giving thirty days written notice.</w:t></w:r></w:p>'
            '</w:sdtContent></w:sdt>'))

    document = parse_document(_docx(build), "sdt.docx")
    assert "TERMINATION" in document.text
    assert "UNILATERAL_TERMINATION" in fired(segment_clauses(document))


def test_tracked_change_insertions_are_read_and_declared():
    """Reviewing a redlined contract is the whole point, and the review was
    reading the version from before the other side's edits."""
    def build(document):
        _base(document)
        document.add_paragraph("3. TERMINATION")
        paragraph = document.add_paragraph("Either party may terminate on thirty days written notice.")
        paragraph._p.append(parse_xml(
            f'<w:ins {nsdecls("w")} w:id="1" w:author="Other Side" w:date="2025-01-01T00:00:00Z">'
            '<w:r><w:t xml:space="preserve"> Notwithstanding the foregoing, the Company '
            'may terminate this Agreement at any time without assigning any reason.</w:t></w:r></w:ins>'))

    document = parse_document(_docx(build), "ins.docx")
    assert "Notwithstanding" in document.text
    assert "UNILATERAL_TERMINATION" in fired(segment_clauses(document))
    assert any("tracked changes" in w for w in document.extraction_warnings)


def test_a_nested_table_is_read():
    def build(document):
        _base(document)
        document.add_paragraph("SCHEDULE")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Item"
        table.cell(0, 1).text = "Amount"
        table.cell(1, 0).text = "Late fee"
        table.cell(1, 1).add_table(rows=1, cols=1).cell(0, 0).text = (
            "a penalty of Rs. 2,000 per day of delay in payment"
        )

    document = parse_document(_docx(build), "nested.docx")
    assert "Rs. 2,000 per day" in document.text


def test_two_cells_that_read_the_same_are_both_kept():
    """Dedupe compared text, so a genuinely repeated amount lost a column."""
    def build(document):
        _base(document)
        document.add_paragraph("SCHEDULE")
        table = document.add_table(rows=1, cols=3)
        table.cell(0, 0).text = "Milestone"
        table.cell(0, 1).text = "Rs. 5,000"
        table.cell(0, 2).text = "Rs. 5,000"

    assert parse_document(_docx(build), "dupe.docx").text.count("Rs. 5,000") == 2


def test_word_automatic_numbering_is_recovered():
    """The numbers live in w:numPr, not in the text, so every finding cited
    "paragraph 7" instead of the clause the reader can actually see."""
    document = docx.Document()
    for heading, body in [
        ("APPOINTMENT", "The Company appoints the Employee with effect from the date stated in the schedule hereto."),
        ("REMUNERATION", "The Employee is paid the salary stated in the schedule monthly in arrears after deductions."),
        ("TERMINATION", "The Company may terminate this Agreement at any time without assigning any reason on thirty days notice."),
        ("NON-COMPETITION", "The Employee shall not join a competing business for twenty four months after the termination of employment."),
    ]:
        document.add_paragraph(heading, style="List Number")
        document.add_paragraph(body)
    buffer = io.BytesIO()
    document.save(buffer)

    _, clauses = parse_document(buffer.getvalue(), "auto.docx"), None
    parsed = parse_document(buffer.getvalue(), "auto.docx")
    clauses = segment_clauses(parsed)
    assert [c.number for c in clauses] == ["1", "2", "3", "4"]
    assert "NON_COMPETE_POST_TERM" in fired(clauses)


# ---------------------------------------------------------------------------
# Encoding and scans
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("encoding", ["utf-16", "utf-16-le", "utf-16-be", "utf-8-sig"])
def test_a_utf16_upload_is_decoded(encoding):
    """Decoded as UTF-8 it became one clause of NUL-separated gibberish that
    every rule silently failed to match."""
    raw = open("tests/fixtures/employment_agreement.txt").read()
    document = parse_document(raw.encode(encoding), "u.txt")
    assert "\x00" not in document.text
    assert len(segment_clauses(document)) > 5
    assert "NON_COMPETE_POST_TERM" in fired(segment_clauses(document))


def _pdf_reader(good_pages, blank_pages, body):
    class Page:
        def __init__(self, text):
            self._text = text

        def extract_text(self):
            return self._text

    pages = [Page(body)] * good_pages + [Page("")] * blank_pages

    class Reader:
        is_encrypted = False

        def __init__(self, *args, **kwargs):
            self.pages = pages

    return Reader


def test_a_mostly_unreadable_pdf_is_refused():
    """Averaged density hid a partial scan: 8 readable pages among 20 pass it
    comfortably while 12 pages of the contract were never read at all."""
    body = open("tests/fixtures/employment_agreement.txt").read()
    with patch("pypdf.PdfReader", _pdf_reader(8, 12, body)):
        with pytest.raises(DocumentParseError, match="scan or an image"):
            parse_document(b"%PDF-1.4 stub", "p.pdf")


def test_a_few_unreadable_pages_are_declared():
    body = open("tests/fixtures/employment_agreement.txt").read()
    with patch("pypdf.PdfReader", _pdf_reader(18, 2, body)):
        document = parse_document(b"%PDF-1.4 stub", "p.pdf")
    assert any("no readable text" in w for w in document.extraction_warnings)


# ---------------------------------------------------------------------------
# Graph nodes: a batch must only write about clauses it was shown
# ---------------------------------------------------------------------------

def _fake_llm(result):
    class Fake:
        def with_structured_output(self, *args):
            return self

        async def ainvoke(self, *args):
            return result

    return Fake()


def _employment():
    document = parse_document(
        open("tests/fixtures/employment_agreement.txt", "rb").read(), "e.txt"
    )
    return document, segment_clauses(document)


def test_a_declared_side_survives_a_guessed_contract_type():
    """The reviewer said which side they were on; a model-guessed contract type
    threw that away and then asked them the same question again."""
    from src.agents.contract_nodes import profiler
    from src.schemas.contract import ContractType
    from src.schemas.contract_review import ContractProfile

    document, clauses = _employment()
    state = {
        "task_id": "t", "document": document, "document_text": document.text,
        "clauses": clauses, "declared_contract_type": None,
        "declared_position": PartyPosition.EMPLOYEE,
        "llm_available": True, "llm_calls_used": 0,
    }
    guessed = ContractProfile(contract_type=ContractType.SERVICE, contract_type_confidence=0.55)
    with patch.object(profiler, "ChatGroq", return_value=_fake_llm(guessed)):
        result = asyncio.run(profiler.run_contract_profiler(state))

    assert result["position"] == PartyPosition.EMPLOYEE.value
    assert result["position_source"] == "USER_DECLARED"
    assert result["contract_type"] == ContractType.EMPLOYMENT.value


def test_an_explainer_batch_cannot_write_about_a_clause_it_never_saw():
    from src.agents.contract_nodes import explainer
    from src.core.clause_triage import triage_clauses
    from src.schemas.contract_review import ClauseExplanation, ClauseExplanationBatch

    document, clauses = _employment()
    findings = evaluate_clauses(clauses, PartyPosition.EMPLOYEE)
    triage = triage_clauses(clauses, findings, PartyPosition.EMPLOYEE)
    victim = triage.hot[0]

    # Only WARM clauses are sent; the batch replies about a HOT one anyway.
    state = {
        "clauses": clauses, "clause_tier": triage.tiers, "hot_clause_indices": [],
        "warm_clause_indices": triage.warm, "position": "EMPLOYEE",
        "clause_chunks": {}, "llm_available": True, "llm_calls_used": 0,
        "degraded_nodes": [],
    }
    batch = ClauseExplanationBatch(explanations=[ClauseExplanation(
        clause_index=victim, plain_english="text about a clause it never saw",
        obligations=[], watch_outs=["invented"])])
    with patch.object(explainer, "ChatGroq", return_value=_fake_llm(batch)):
        result = asyncio.run(explainer.run_clause_explainer(state))

    assert "never saw" not in result["explanations"][victim]["plain_english"]


def test_a_scoped_retry_ignores_findings_about_other_clauses():
    """The retry is scoped to the clauses whose quotes failed to ground. A
    finding about a different clause duplicated one the first pass had made."""
    from src.agents.contract_nodes import red_flag_analyst
    from src.core.clause_triage import triage_clauses
    from src.schemas.contract_review import LlmRedFlag, LlmRedFlagBatch

    document, clauses = _employment()
    findings = evaluate_clauses(clauses, PartyPosition.EMPLOYEE)
    triage = triage_clauses(clauses, findings, PartyPosition.EMPLOYEE)
    scoped, outside = triage.hot[0], triage.hot[1]

    state = {
        "clauses": clauses, "position": "EMPLOYEE", "rule_findings": findings,
        "hot_clause_indices": triage.hot, "ungrounded_clause_indices": [scoped],
        "retry_count": 1, "llm_findings": [], "suppressed_llm_findings": [],
        "degraded_nodes": [], "llm_available": True, "llm_calls_used": 0,
    }
    batch = LlmRedFlagBatch(findings=[LlmRedFlag(
        clause_index=outside, title="Out of batch",
        quote=clauses[outside].text[:60], plain_summary="s", why_it_matters="w")])
    with patch.object(red_flag_analyst, "ChatGroq", return_value=_fake_llm(batch)):
        result = asyncio.run(red_flag_analyst.run_red_flag_analyst(state))

    assert outside not in [f.clause_index for f in result["llm_findings"]]
    assert any("not in this batch" in s["reason"] for s in result["suppressed_llm_findings"])


def test_two_findings_with_similar_titles_are_both_kept():
    """Keyed on the first 40 characters of the title alone, "... brought by
    third parties" and "... brought by the Company" collapsed into one and the
    second was lost with nothing recorded."""
    from src.agents.contract_nodes import red_flag_analyst
    from src.schemas.contract_review import LlmRedFlag, LlmRedFlagBatch

    document, clauses = _employment()
    target = next(c for c in clauses if len(c.text) > 200)
    batch = LlmRedFlagBatch(findings=[
        LlmRedFlag(clause_index=target.index,
                   title="Employee bears cost of all litigation brought by third parties",
                   quote=target.text[:50], plain_summary="A", why_it_matters="w"),
        LlmRedFlag(clause_index=target.index,
                   title="Employee bears cost of all litigation brought by the Company itself",
                   quote=target.text[60:110], plain_summary="B", why_it_matters="w"),
    ])
    state = {
        "clauses": clauses, "position": "EMPLOYEE", "rule_findings": [],
        "hot_clause_indices": [target.index], "llm_findings": [],
        "suppressed_llm_findings": [], "degraded_nodes": [],
        "llm_available": True, "llm_calls_used": 0, "retry_count": 0,
    }
    with patch.object(red_flag_analyst, "ChatGroq", return_value=_fake_llm(batch)):
        result = asyncio.run(red_flag_analyst.run_red_flag_analyst(state))

    assert len(result["llm_findings"]) == 2
    assert verify_quotes(result["llm_findings"], document.text) == []


def test_the_classifier_does_not_relabel_the_callers_clauses():
    """It mutated the caller's Clause objects, so the deterministic categories
    it recorded were already the model's, and a re-run saw different input."""
    from src.agents.contract_nodes import clause_classifier
    from src.schemas.contract import ClauseCategory
    from src.schemas.contract_review import ClauseLabel, ClauseLabelBatch

    _, clauses = _employment()
    target = next(c for c in clauses if c.category is ClauseCategory.NON_COMPETE)
    before = target.category

    batch = ClauseLabelBatch(labels=[ClauseLabel(
        clause_index=target.index, category=ClauseCategory.DEFINITIONS,
        confidence=0.95, reason="x")])
    state = {
        "clauses": clauses, "position": "EMPLOYEE",
        "llm_available": True, "llm_calls_used": 0, "degraded_nodes": [],
    }
    with patch.object(clause_classifier, "ChatGroq", return_value=_fake_llm(batch)):
        result = asyncio.run(clause_classifier.run_clause_classifier(state))

    assert target.category is before, "caller's clause list was mutated"
    assert result["deterministic_categories"][target.index] == before.value
